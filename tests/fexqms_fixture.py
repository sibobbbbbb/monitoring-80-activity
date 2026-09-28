"""Pembangun file export FEXQMS sintetis untuk test, dengan layout sama seperti export asli:
blok header (label -> nilai, USL/LSL di header), baris judul tabel di baris 9, data mulai baris 10."""
import csv
import io
from datetime import datetime, timedelta

import openpyxl

TABLE_HEADER = [
    "No.", "Sample ID", "Sample Date Time", "Part", None, "Machine", None, "Characteristics", "Lot",
    None, None, "Inspected By", "Shift", None, "Mea. Data : 1", None, "XChart", "MovingRange",
    "Violation (XChart)", "Violation (MovingRange)", "Sample Remarks", "Root Cause Remarks",
    "Action Taken Remarks", "Close Remarks",
]
WIDTH = len(TABLE_HEADER)
FIRST_DATA_ROW = 10   # nomor baris Excel baris data pertama (judul tabel di baris 9)

PART, OPERATION, MACHINE, CHARACTERISTICS = "Demo Part", "Demo Op 10", "MC-1", "Demo Flatness 1"
T0 = datetime(2026, 1, 5, 8, 0, 0)


def _row(*cells) -> list:
    return list(cells) + [None] * (WIDTH - len(cells))


def times(n: int, start: datetime = T0, step_minutes: int = 30) -> list[datetime]:
    return [start + timedelta(minutes=step_minutes * i) for i in range(n)]


def export_grid(samples, part=PART, operation=OPERATION, machine=MACHINE, characteristics=CHARACTERISTICS,
                usl=0.2, lsl=0, nominal=None, per_row_identity=True) -> list[list]:
    """samples: list of (waktu, xchart). Kolom 'Mea. Data : 1' diisi 2x XChart agar terbukti tidak dipakai."""
    grid = [
        _row(),
        _row(None, None, None, None, "Statistical Process Control Chart"),
        _row("Part:", None, part, None, "Operation:", None, operation, None, None, None, "Machine:", machine),
        _row("Characteristics:", None, characteristics, None, "Measurement Unit:", None, "mm", None, None, None,
             "Sampling Frequency:", "Once / 50 pcs"),
        _row("Equipment:", None, "CMM", None, "USL:", None, usl, None, None, None, "LSL:", lsl),
        _row("Control Plan Name:", None, "CP_Demo", None, "Sample ID:", None, 1, None, None, None,
             *(["Nominal:", nominal] if nominal is not None else ["", ""])),
        _row(),
        _row(),
        _row(*TABLE_HEADER),
    ]
    for n, (when, x) in enumerate(samples, 1):
        raw = x * 2 if isinstance(x, (int, float)) else None
        grid.append(_row(
            n, 1000 + n, when,
            part if per_row_identity else None, None,
            machine if per_row_identity else None, None,
            characteristics if per_row_identity else None,
            "L1", None, None, "Inspector", "White Pagi", None, raw, None, x,
        ))
    return grid


def to_xlsx(grid, data_sheet_index: int = 1) -> bytes:
    """Workbook dengan Sheet1 (ringkasan pelanggaran, diabaikan) dan Sheet2 (data), seperti aslinya."""
    wb = openpyxl.Workbook()
    summary = wb.active
    summary.title = "Sheet1"
    summary["H1"] = "Statistical Process Control Chart"
    summary.append([])
    summary.append(["No", "Date & Time", None, "OOC Point No", "OOC Value", "Rule Violation"])
    data = wb.create_sheet("Sheet2")
    if data_sheet_index == 0:                      # data di sheet pertama saja
        wb.remove(summary)
    for r, row in enumerate(grid, 1):
        for c, v in enumerate(row, 1):
            if v is not None:
                data.cell(row=r, column=c, value=v)
    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()


def to_csv(grid, delimiter: str = ",", bom: bool = False) -> bytes:
    buf = io.StringIO()
    writer = csv.writer(buf, delimiter=delimiter, lineterminator="\n")
    for row in grid:
        writer.writerow(["" if v is None else v for v in row])
    text = buf.getvalue()
    return (b"\xef\xbb\xbf" if bom else b"") + text.encode("utf-8")
